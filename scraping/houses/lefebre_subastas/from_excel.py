"""Conversor Excel -> JSONL para Lefebre Subastas.

Lefebre es la primera casa que NO se scrapea: su dato solo existe en una hoja
curada a mano en 2024 (FINALL.xlsx, hoja "FINAL"). Este modulo hace el papel que
en las demas casas hace el scraper: leer la fuente y dejar JSONL de LotItem en
output/. De ahi en adelante Lefebre es indistinguible del resto del pipeline,
porque bronze/ingest.py solo copia *.jsonl y nunca importa el modulo de la casa.

Se ejecuta a mano, una vez, igual que un scraper. NO entra en run_all.ps1:

    python -m scraping.houses.lefebre_subastas.from_excel --excel FINALL.xlsx

La hoja mezcla dos casas: se filtra por 'Casa de Subasta' == 'Lefebre'. Las 4.073
filas de Bogota se descartan a proposito, porque esa casa ya entra por su propio
scraper y mezclarlas duplicaria el dato.
"""

from __future__ import annotations

import argparse
import re
import unicodedata
from datetime import date, datetime
from pathlib import Path
from typing import Any, Optional

import orjson
from openpyxl import load_workbook

from scraping.common.models import LotItem

ROOT = Path(__file__).resolve().parents[3]
DEFAULT_EXCEL = ROOT / "FINALL.xlsx"
DEFAULT_OUTPUT_DIR = ROOT / "scraping" / "houses" / "lefebre_subastas" / "output"

SHEET = "FINAL"
HOUSE_SLUG = "lefebre_subastas"
HOUSE_NAME = "Lefebre Subastas"
HOUSE_VALUE = "Lefebre"  # como aparece en la columna 'Casa de Subasta'
CURRENCY = "COP"

# La hoja usa estos centinelas donde no hay dato. Se tratan como vacio en vez de
# arrastrarlos como texto: '#N/D' es un error de formula de Excel, no un pais.
SENTINELS = {"", "#N/D", "#N/A", "UNKNOWN", "NO IMPORTA NO ARR", "NAN", "NONE"}

# 'Pasado' y '0' significan lo MISMO: el lote no se vendio. La diferencia es de
# convencion por subasta (las 27/28/29/30 escriben 0, las anteriores 'Pasado'),
# no semantica. Confundirlas invertiria la tasa de venta de la casa.
NOT_SOLD_TOKENS = {"PASADO", "SOLD AMOUNT NOT FOUND", "PENDIENTE"}


def _clean(value: Any) -> Optional[str]:
    """Texto util o None. Colapsa los centinelas de la hoja a None."""
    if value is None:
        return None
    text = str(value).strip()
    if not text or text.upper() in SENTINELS:
        return None
    return text


def _to_int(value: Any) -> Optional[int]:
    """Entero o None. Acepta '1.200.000', '1,200,000' y 'COP$1,200,000'.

    Los numeros se atajan ANTES de pasar por texto: openpyxl devuelve las celdas
    numericas como float (1.0, 1250000.0) y el limpiado de separadores de millar
    se comeria el punto decimal, convirtiendo 1.0 en 10 y 1250000.0 en 12500000
    (todo x10). El bug afectaba a precios y a numeros de lote por igual.
    """
    if isinstance(value, bool):  # bool es subclase de int: no es una cifra
        return None
    if isinstance(value, (int, float)):
        return int(value)

    text = _clean(value)
    if text is None:
        return None
    if text.upper() in NOT_SOLD_TOKENS:
        return None
    # Quita moneda y separadores de millar en cualquiera de los dos estilos.
    # Aqui ya solo llegan cadenas: '1.200.000' y '1,200,000' son separadores.
    digits = re.sub(r"[^\d]", "", text)
    if not digits:
        return None
    return int(digits)


def _year(value: Any) -> Optional[int]:
    """Anio de nacimiento/muerte. 'Vive' -> None (no es una fecha)."""
    text = _clean(value)
    if text is None:
        return None
    match = re.search(r"(?:1[6-9]|20)\d{2}", text)
    return int(match.group(0)) if match else None


def slugify(text: str) -> str:
    """'Subasta Barranquilla' -> 'subasta-barranquilla'.

    Sin tildes ni signos: el slug acaba dentro de auction_id y de las URIs
    sinteticas, y esas viajan por todo el pipeline como claves.
    """
    norm = unicodedata.normalize("NFKD", text)
    norm = "".join(c for c in norm if not unicodedata.combining(c))
    norm = re.sub(r"[^a-zA-Z0-9]+", "-", norm).strip("-").lower()
    return norm or "sin-nombre"


def normalize_date(value: Any) -> Optional[str]:
    """Fecha de subasta -> ISO 'YYYY-MM-DD'.

    La columna mezcla datetime reales y texto 'dd/mm/YYYY'. Se normalizan los dos
    a ISO porque extract_year() lee los 4 primeros caracteres: en ISO devuelve
    ("2024", "iso") y no ("2024", "text"), que el informe marcaria como inferido.
    """
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    text = _clean(value)
    if text is None:
        return None
    # 'dd/mm/YYYY' (formato de la hoja, dia primero).
    match = re.match(r"^(\d{1,2})/(\d{1,2})/(\d{4})$", text)
    if match:
        day, month, year = (int(g) for g in match.groups())
        try:
            return date(year, month, day).isoformat()
        except ValueError:
            return None
    # Ya venia en ISO (a veces con hora detras).
    match = re.match(r"^(\d{4}-\d{2}-\d{2})", text)
    if match:
        return match.group(1)
    return None


# 'Medidas:' abre las dimensiones y cierra el resto del bloque.
_MEDIDAS_RE = re.compile(r"medidas\s*:\s*", re.IGNORECASE)
# Parentesis biografico del artista: '(1907-1992)', '(1957)', '( 1911 - 2022)'.
# Marca donde acaba el nombre y empieza el titulo de la obra.
_BIO_RE = re.compile(r"\(\s*\d{4}\s*(?:[-–]\s*\d{0,4}\s*)?\)")
# La tecnica es un vocabulario cerrado y corto. Sirve de corte cuando el titulo
# viene pegado a ella con un solo espacio ('S/T, 1974Tecnica mixta...'), que es
# lo unico que no distingue la separacion por 2+ espacios.
_MEDIUM_RE = re.compile(
    r"(óleo|oleo|acrílico|acrilico|acuarela|gouache|témpera|tempera|pastel|"
    r"tinta|carboncillo|grabado|litografía|litografia|serigrafía|serigrafia|"
    r"aguafuerte|xilografía|xilografia|obra gráfica|obra grafica|fotografía|"
    r"fotografia|técnica mixta|tecnica mixta|escultura|bronce|mármol|marmol|"
    r"madera|cerámica|ceramica|ensamble|collage|dibujo|díptico|diptico)",
    re.IGNORECASE,
)


def split_title(raw: Optional[str]) -> tuple[Optional[str], Optional[str], Optional[str]]:
    """Parte el bloque de 'Title' en (titulo, tecnica, medidas).

    En Lefebre las columnas Tecnica/Tamanio estan vacias en las 1.711 filas: lo
    unico que hay es este bloque, con la forma

        ARTISTA (1907-1992)
        Titulo de la obra, 1984
        Oleo sobre masonite
        Medidas: 70 x 120 cm

    OJO: las 4 subastas mas nuevas (27/28/29/30, 678 lotes) traen exactamente lo
    mismo pero con ESPACIOS en vez de saltos de linea. Partir solo por lineas
    dejaba esos lotes sin titulo. Por eso se corta por dos marcas que existen en
    los dos formatos: el parentesis biografico y 'Medidas:'.
    """
    if not raw:
        return None, None, None
    text = str(raw).strip()
    if not text:
        return None, None, None

    # 1. Las dimensiones van detras de 'Medidas:'; el resto queda delante.
    dimensions = None
    parts = _MEDIDAS_RE.split(text, maxsplit=1)
    head = parts[0]
    if len(parts) > 1:
        # Corta la edicion ('Ed. 74/100') si viene pegada detras de las medidas.
        dimensions = re.split(r"\s{2,}|\n|\bEd\.", parts[1], maxsplit=1)[0].strip() or None

    # 2. Quita el nombre del artista: va delante del parentesis biografico.
    bio = _BIO_RE.search(head)
    if bio:
        rest = head[bio.end():]
    elif "\n" in head:
        # Sin parentesis pero multilinea: la 1a linea sigue siendo el artista.
        rest = "\n".join(head.splitlines()[1:])
    else:
        # Sin parentesis y en una sola linea no hay forma fiable de separar el
        # nombre del titulo, asi que se deja entero antes que cortar por donde
        # no toca. Artista_Clean ya trae el nombre por su cuenta.
        rest = head

    # 3. Lo que queda es 'titulo / tecnica': separados por salto de linea o por
    #    2+ espacios, y en las subastas nuevas a veces por nada en absoluto.
    chunks = [c.strip(" ,;") for c in re.split(r"\n|\s{2,}", rest) if c.strip(" ,;")]
    title = chunks[0] if chunks else None
    medium = chunks[1] if len(chunks) > 1 else None

    # Titulo y tecnica pegados: se corta por el nombre de la tecnica.
    if title and medium is None:
        found = _MEDIUM_RE.search(title)
        if found and found.start() > 0:
            medium = title[found.start():].strip(" ,;")
            title = title[: found.start()].strip(" ,;")

    return title or None, medium, dimensions


def real_lot_numbers(rows: list[dict]) -> bool:
    """True si 'Order' es de verdad el numero de lote de esta subasta.

    Solo las subastas 27/28/29/30 traen numeracion real (1..N contigua). En las
    11 antiguas 'Order' es el numero de FILA de la hoja (~4074-5110). Volcarlo
    como lot_number fabricaria huecos enormes en la secuencia, y build_gold.py
    lee esos huecos para avisar del sesgo de tasa de venta: saldria un aviso
    inventado. Por eso ahi lot_number se queda a None.
    """
    orders = [r["order"] for r in rows if r["order"] is not None]
    if not orders:
        return False
    return min(orders) == 1 and max(orders) == len(orders) == len(set(orders))


def read_rows(excel_path: Path) -> list[dict]:
    """Lee la hoja FINAL y devuelve solo las filas de Lefebre, ya tipadas."""
    workbook = load_workbook(excel_path, read_only=True, data_only=True)
    if SHEET not in workbook.sheetnames:
        raise SystemExit(f"La hoja '{SHEET}' no existe en {excel_path}")
    sheet = workbook[SHEET]

    stream = sheet.iter_rows(values_only=True)
    header = [str(c).strip() if c is not None else "" for c in next(stream)]
    index = {name: i for i, name in enumerate(header)}

    def cell(record: tuple, column: str) -> Any:
        position = index.get(column)
        return record[position] if position is not None and position < len(record) else None

    rows = []
    for record in stream:
        if not record or _clean(cell(record, "Casa de Subasta")) != HOUSE_VALUE:
            continue
        auction_name = _clean(cell(record, "Auction Name")) or "Sin subasta"
        order = _to_int(cell(record, "Order"))
        rows.append(
            {
                "auction_name": auction_name,
                "order": order,
                "title_raw": _clean(cell(record, "Title")),
                "title_clean": _clean(cell(record, "Título_Clean")),
                "lot_url": _clean(cell(record, "URL Lote")),
                "auction_url": _clean(cell(record, "URL Subasta")),
                "date": normalize_date(cell(record, "Auction_Date")),
                "estimate": _to_int(cell(record, "Starting Price_Clean")),
                # Se guarda el valor CRUDO de la celda, sin pasar por _clean():
                # convertirlo a texto aqui haria que un float 900000.0 llegase a
                # _to_int como '900000.0' y el punto se leyese como separador de
                # millar (x10). El valor se interpreta en build_lot().
                "sold_raw": cell(record, "Sold For_Clean"),
                "artist": _clean(cell(record, "Artista_Clean")),
                "artist_raw": _clean(cell(record, "Artista_Raw")),
                "country": _clean(cell(record, "Pais_Clean")),
                "birth": _year(cell(record, "Nacimiento_Clean")),
                "death": _year(cell(record, "Muerte_Clean")),
            }
        )
    workbook.close()
    return rows


def build_lot(row: dict, auction_slug: str, use_lot_number: bool) -> LotItem:
    """Construye un LotItem a partir de una fila ya tipada.

    Se usa el modelo compartido y no un dict a mano porque model_dump() garantiza
    las 31 claves que Gold espera leer con .get(): una clave ausente se convierte
    en un None silencioso aguas abajo.
    """
    # Precio: solo cuenta si es un numero > 0. 'Pasado' y 0 son no vendido.
    sold = _to_int(row["sold_raw"])
    if sold is not None and sold <= 0:
        sold = None
    status = "VENDIDO" if sold is not None else "NO VENDIDO"

    # Silver descarta toda fila sin lot_url y deduplica por house|lot_url, asi
    # que el id tiene que ser unico Y estable entre ejecuciones: si cambiara,
    # cada ingesta duplicaria los lotes. (auction_name, order) no colisiona.
    lot_url = row["lot_url"] or f"excel://{HOUSE_SLUG}/{auction_slug}/{row['order']}"

    auction_url = row["auction_url"]
    if not auction_url or not auction_url.lower().startswith("http"):
        # En Lefebre esta columna trae el nombre de la subasta, no una URL.
        auction_url = f"excel://{HOUSE_SLUG}/{auction_slug}"

    title, medium, dimensions = split_title(row["title_raw"])

    return LotItem(
        auction_id=auction_slug,
        auction_title=row["auction_name"],
        auction_start_date=row["date"],
        auction_house_name=HOUSE_NAME,
        auction_url=auction_url,
        lot_number=row["order"] if use_lot_number else None,
        lot_url=lot_url,
        lot_title=row["title_clean"] or title,
        price_estimate_min=row["estimate"],
        price_estimate_max=None,  # la hoja solo tiene precio de salida
        price_sold=sold,
        currency=CURRENCY,
        artist_name=row["artist"],
        artist_birth_year=row["birth"],
        artist_death_year=row["death"],
        # Diagnostico unicamente: alimenta unmapped_country_values y el contador
        # de conflictos. El pais real lo fija el maestro en artist_resolve.py,
        # nunca el texto libre de la casa.
        artist_country=row["country"],
        artist_raw=row["artist_raw"],
        description=row["title_raw"],
        medium=medium,
        dimensions=dimensions,
        status=status,
    )


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--excel", type=Path, default=DEFAULT_EXCEL)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)
    args = parser.parse_args()

    if not args.excel.exists():
        raise SystemExit(f"No existe el Excel: {args.excel}")

    rows = read_rows(args.excel)
    if not rows:
        raise SystemExit(f"Ninguna fila con 'Casa de Subasta' == {HOUSE_VALUE!r}")

    by_auction: dict[str, list[dict]] = {}
    for row in rows:
        by_auction.setdefault(row["auction_name"], []).append(row)

    args.output_dir.mkdir(parents=True, exist_ok=True)
    total = sold_total = 0

    for auction_name, auction_rows in sorted(by_auction.items()):
        auction_slug = slugify(auction_name)
        use_lot_number = real_lot_numbers(auction_rows)

        path = args.output_dir / f"{auction_slug}.jsonl"
        with open(path, "wb") as handle:
            for row in auction_rows:
                lot = build_lot(row, auction_slug, use_lot_number)
                handle.write(orjson.dumps(lot.model_dump(), option=orjson.OPT_APPEND_NEWLINE))
                total += 1
                sold_total += lot.price_sold is not None

        numbering = "lot_number real" if use_lot_number else "sin lot_number"
        print(f"[lefebre] {auction_slug}: {len(auction_rows):4d} lotes ({numbering})")

    print(f"[lefebre] {total} lotes en {len(by_auction)} subastas -> {args.output_dir}")
    print(f"[lefebre] vendidos: {sold_total} ({sold_total / total:.1%})")


if __name__ == "__main__":
    main()
