#!/usr/bin/env python3
"""Construye el maestro operativo exhaustivo y su registro de cobertura.

El YAML de ``pipelines/config/artists`` sigue siendo la fuente curada de
identidades. Este modulo NO convierte automaticamente cada texto de catalogo
en una persona: combina las identidades curadas con todas las observaciones de
autor de Silver y separa de forma explicita cuatro estados:

``verified``
    Identidad curada, con evidencia URL y sin huecos o conflictos detectados.
``needs_web``
    Identidad/candidato que necesita completar o documentar campos.
``needs_review``
    Identidad ya curada, pero con conflicto de catalogo o evidencia insuficiente.
``ambiguous``
    Observacion sin identidad estable (homonimo, posible alias o datos en conflicto).

De esta forma el 100% significa "contabilizado y clasificado", no "promovido
sin pruebas". Los valores de catalogo viajan como candidatos y nunca pisan
``country_birth`` ni las fechas canonicas del maestro.

Salidas locales (gitignored):

* ``data/gold/dim_artist_master.jsonl``: maestro operativo combinado.
* ``data/gold/artist_name_disposition.jsonl``: disposicion de cada fold observado.
* ``data/gold/artist_master_coverage.json``: prueba de cobertura e indicadores.

Con ``--write-review-queue`` se materializa ademas la cola versionada y
fragmentada en ``pipelines/config/artist_review/*.jsonl``. Esos ficheros
pertenecen al proceso de curaduria, pero ``load_master()`` no los carga: una
fila pendiente no adquiere identidad ni pais por el mero hecho de estar
registrada.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable

from pipelines.shared.artist_key import (
    artist_fold,
    attribution_type,
    propose_inversion,
    strip_biography,
)
from pipelines.shared.artist_master import (
    load_master,
    normalize_country,
    resolve_artist,
)


ROOT = Path(__file__).resolve().parents[2]
SILVER_LOTS = ROOT / "data" / "silver" / "lots.jsonl"
GOLD_ROOT = ROOT / "data" / "gold"
DIM_ARTISTS = GOLD_ROOT / "dim_artist_master.jsonl"
DISPOSITIONS = GOLD_ROOT / "artist_name_disposition.jsonl"
COVERAGE = GOLD_ROOT / "artist_master_coverage.json"
LEGACY_EXCEL = ROOT / "FINALL.xlsx"
REVIEW_DIR = ROOT / "pipelines" / "config" / "artist_review"
REVIEW_MANIFEST = REVIEW_DIR / "_manifest.json"

CURATION_STATUSES = ("verified", "needs_web", "needs_review", "ambiguous")
CANONICAL_FIELDS = ("country_birth", "birth_year", "death_year")

_PAREN_RE = re.compile(r"\(([^()]*)\)")
_YEAR_RE = re.compile(r"\b(1[0-9]{3}|20[0-9]{2})\b")
_URL_RE = re.compile(r"https?://[^\s;|]+")


def _json_value(value: Any) -> Any:
    """Convierte claves de Counter a valores JSON estables."""
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    return str(value)


def _safe_year(value: Any) -> int | None:
    """Anio plausible o None; nunca interpreta centinelas como fechas."""
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, datetime):
        value = value.year
    try:
        year = int(float(value))
    except (TypeError, ValueError):
        return None
    # El corpus contiene maestros antiguos; el limite inferior solo elimina
    # seriales/corrupcion obvia, no impone que el arte sea moderno.
    if 500 <= year <= date.today().year:
        return year
    return None


def _counter_rows(counter: Counter) -> list[dict[str, Any]]:
    return [
        {"value": _json_value(value), "count": count}
        for value, count in sorted(
            counter.items(), key=lambda item: (-item[1], str(item[0]))
        )
    ]


def _single(counter: Counter) -> Any:
    return next(iter(counter)) if len(counter) == 1 else None


def _notes_urls(notes: Any) -> list[str]:
    if not notes:
        return []
    return sorted({url.rstrip(".,)") for url in _URL_RE.findall(str(notes))})


def parse_biography_evidence(text: Any) -> dict[str, Any]:
    """Extrae solo senales explicitas de parentesis biograficos.

    El resultado sigue siendo evidencia de catalogo, no verdad maestra. La
    primera localidad/pais y los dos primeros anios se guardan como candidatos
    para investigacion posterior.
    """
    if not text:
        return {}
    result: dict[str, Any] = {}
    for chunk in _PAREN_RE.findall(str(text)):
        years = [_safe_year(value) for value in _YEAR_RE.findall(chunk)]
        years = [value for value in years if value is not None]
        if years:
            result.setdefault("birth_year", years[0])
            if len(years) > 1:
                result.setdefault("death_year", years[1])

        head = re.split(r"[,;]", chunk, maxsplit=1)[0].strip()
        if head and not _YEAR_RE.search(head):
            code = normalize_country(head)
            if code:
                result.setdefault("country_birth", code)
    return result


def _new_group(key: str, artist_id: str | None) -> dict[str, Any]:
    return {
        "artist_key": key,
        "artist_id": artist_id,
        "names": Counter(),
        "folds": set(),
        "houses": Counter(),
        "sources": Counter(),
        "source_files": set(),
        "lot_urls": [],
        "lots_observed": 0,
        "legacy_rows": 0,
        "catalog_countries": Counter(),
        "bio_countries": Counter(),
        "birth_years": Counter(),
        "death_years": Counter(),
    }


def _append_sample(values: list[str], value: Any, limit: int = 5) -> None:
    if value and value not in values and len(values) < limit:
        values.append(str(value))


def _resolved_identity(row: dict[str, Any], name: str) -> tuple[str | None, str | None, str]:
    """Devuelve (artist_id, fold, resolution), tolerando Silver antiguo."""
    resolution = row.get("artist_resolution")
    fold = row.get("artist_fold") or artist_fold(name)
    artist_id = row.get("artist_id")
    if resolution is None:
        resolved = resolve_artist(name)
        resolution = resolved["artist_resolution"]
        fold = resolved["artist_fold"] or fold
        artist_id = resolved["artist_id"]
    return artist_id, fold, resolution or "fold_only"


def _add_author_observation(
    groups: dict[str, dict[str, Any]],
    row: dict[str, Any],
    *,
    source: str,
    count_lot: bool,
) -> tuple[str | None, str | None]:
    name = str(row.get("artist_name") or "").strip()
    if not name:
        return None, None
    artist_id, fold, _resolution = _resolved_identity(row, name)
    if not fold and not artist_id:
        return None, None

    key = artist_id or fold
    group = groups.setdefault(key, _new_group(key, artist_id))
    group["names"][name] += 1
    if fold:
        group["folds"].add(fold)
    if row.get("house_slug"):
        group["houses"][str(row["house_slug"])] += 1
    group["sources"][source] += 1
    if row.get("source_file"):
        group["source_files"].add(str(row["source_file"]))
    _append_sample(group["lot_urls"], row.get("lot_url"))
    if count_lot:
        group["lots_observed"] += 1
    else:
        group["legacy_rows"] += 1

    country = normalize_country(row.get("artist_country"))
    if country:
        group["catalog_countries"][country] += 1

    birth = _safe_year(row.get("artist_birth_year"))
    death = _safe_year(row.get("artist_death_year"))
    if birth:
        group["birth_years"][birth] += 1
    if death:
        group["death_years"][death] += 1

    # Se inspeccionan ambos campos, pero un mismo valor solo suma una vez por
    # fila y tipo de evidencia para no premiar una repeticion del scraper.
    bio_values: dict[str, set[Any]] = defaultdict(set)
    for text in (name, row.get("artist_raw")):
        for field, value in parse_biography_evidence(text).items():
            bio_values[field].add(value)
    for value in bio_values["country_birth"]:
        group["bio_countries"][value] += 1
    for value in bio_values["birth_year"]:
        group["birth_years"][value] += 1
    for value in bio_values["death_year"]:
        group["death_years"][value] += 1
    return key, fold


def _new_disposition(fold: str | None, attribution: str) -> dict[str, Any]:
    return {
        "artist_fold": fold,
        "attribution_type": attribution,
        "names": Counter(),
        "lots": 0,
        "houses": Counter(),
        "artist_key": None,
        "sources": Counter(),
    }


def _disposition_key(fold: str | None, attribution: str) -> str:
    return f"{attribution}:{fold or '<none>'}"


def collect_silver(
    lots_path: Path,
) -> tuple[dict[str, dict[str, Any]], dict[str, dict[str, Any]], list[dict[str, Any]], dict[str, int]]:
    if not lots_path.exists():
        raise FileNotFoundError(
            f"Silver no encontrado: {lots_path}. Ejecuta antes artist_resolve."
        )

    groups: dict[str, dict[str, Any]] = {}
    dispositions: dict[str, dict[str, Any]] = {}
    latent: dict[str, dict[str, Any]] = {}
    stats = Counter()

    with open(lots_path, encoding="utf-8") as handle:
        for line in handle:
            line = line.strip()
            if not line:
                continue
            row = json.loads(line)
            stats["silver_lots"] += 1
            name = str(row.get("artist_name") or "").strip()
            attribution = row.get("attribution_type") or attribution_type(name)
            fold = row.get("artist_fold") or artist_fold(name)

            dkey = _disposition_key(fold, attribution)
            disp = dispositions.setdefault(dkey, _new_disposition(fold, attribution))
            if name:
                disp["names"][name] += 1
            disp["lots"] += 1
            if row.get("house_slug"):
                disp["houses"][str(row["house_slug"])] += 1
            disp["sources"]["silver"] += 1

            if attribution == "autor" and fold:
                stats["author_lots"] += 1
                key, _ = _add_author_observation(
                    groups, row, source="silver", count_lot=True
                )
                disp["artist_key"] = key
                continue

            stats[f"attribution_{attribution}"] += 1

            # Cola conservadora: solo primera linea de Lefebre cuando no existe
            # nombre estructurado. Se registra aparte y nunca se promueve a la
            # dimension como artista confirmado.
            if (
                attribution == "desconocido"
                and row.get("house_slug") == "lefebre_subastas"
                and row.get("description")
            ):
                lines = [
                    part.strip()
                    for part in str(row["description"]).splitlines()
                    if part.strip()
                ]
                if len(lines) >= 2:
                    candidate = strip_biography(lines[0])
                    candidate_fold = artist_fold(candidate)
                    if candidate_fold and attribution_type(candidate) == "autor":
                        item = latent.setdefault(
                            candidate_fold,
                            {
                                "artist_fold": candidate_fold,
                                "display_name": candidate,
                                "raw_names": Counter(),
                                "lots": 0,
                                "sample_lot_urls": [],
                                "status": "needs_manual_extraction",
                            },
                        )
                        item["raw_names"][candidate] += 1
                        item["lots"] += 1
                        _append_sample(item["sample_lot_urls"], row.get("lot_url"))

    latent_rows = []
    for item in latent.values():
        item = dict(item)
        item["raw_names"] = [name for name, _ in item["raw_names"].most_common()]
        latent_rows.append(item)
    latent_rows.sort(key=lambda row: (-row["lots"], row["artist_fold"]))
    return groups, dispositions, latent_rows, dict(stats)


def _excel_row(
    name: Any,
    country: Any,
    birth: Any,
    death: Any,
    *,
    source_file: str,
    lot_url: Any = None,
) -> dict[str, Any] | None:
    if name is None or not str(name).strip():
        return None
    text = str(name).strip()
    resolved = resolve_artist(text)
    return {
        "artist_name": text,
        "artist_raw": text,
        "artist_country": country,
        "artist_birth_year": birth,
        "artist_death_year": death,
        "artist_id": resolved["artist_id"],
        "artist_fold": resolved["artist_fold"],
        "artist_resolution": resolved["artist_resolution"],
        "attribution_type": resolved["attribution_type"],
        "house_slug": "legacy_finall",
        "source_file": source_file,
        "lot_url": lot_url,
    }


def collect_legacy_excel(
    excel_path: Path,
    groups: dict[str, dict[str, Any]],
    dispositions: dict[str, dict[str, Any]],
) -> dict[str, int]:
    """Incorpora observaciones explicitas de las hojas Artistas y FINAL.

    Es opcional porque FINALL.xlsx esta gitignored. Sus valores son evidencia
    candidata; nunca sobrescriben el YAML curado.
    """
    stats = Counter()
    if not excel_path.exists():
        return {"legacy_excel_present": 0}

    try:
        from openpyxl import load_workbook
    except ImportError:
        return {"legacy_excel_present": 1, "legacy_excel_skipped_no_openpyxl": 1}

    workbook = load_workbook(excel_path, read_only=False, data_only=True)
    stats["legacy_excel_present"] = 1

    def consume(row: dict[str, Any], source: str) -> None:
        name = str(row.get("artist_name") or "").strip()
        attribution = row.get("attribution_type") or attribution_type(name)
        fold = row.get("artist_fold") or artist_fold(name)
        dkey = _disposition_key(fold, attribution)
        disp = dispositions.setdefault(dkey, _new_disposition(fold, attribution))
        if name:
            disp["names"][name] += 1
        disp["sources"][source] += 1
        if attribution == "autor" and fold:
            key, _ = _add_author_observation(
                groups, row, source=source, count_lot=False
            )
            disp["artist_key"] = key
            stats["legacy_author_rows"] += 1
        else:
            stats["legacy_excluded_rows"] += 1

    if "Artistas" in workbook.sheetnames:
        sheet = workbook["Artistas"]
        for values in sheet.iter_rows(min_row=2, values_only=True):
            if not values:
                continue
            padded = tuple(values) + (None,) * max(0, 4 - len(values))
            row = _excel_row(
                padded[0],
                padded[1],
                padded[2],
                padded[3],
                source_file="FINALL.xlsx#Artistas",
            )
            if row:
                consume(row, "legacy_artistas")

    if "FINAL" in workbook.sheetnames:
        sheet = workbook["FINAL"]
        iterator = sheet.iter_rows(values_only=True)
        header = next(iterator, ())
        index = {
            str(value).strip(): position
            for position, value in enumerate(header)
            if value is not None
        }

        def cell(values: tuple[Any, ...], column: str) -> Any:
            position = index.get(column)
            return values[position] if position is not None and position < len(values) else None

        for values in iterator:
            if not values:
                continue
            row = _excel_row(
                cell(values, "Artista_Clean"),
                cell(values, "Pais_Clean"),
                cell(values, "Nacimiento_Clean"),
                cell(values, "Muerte_Clean"),
                source_file="FINALL.xlsx#FINAL",
                lot_url=cell(values, "URL Lote"),
            )
            if row:
                consume(row, "legacy_final")

    workbook.close()
    return dict(stats)


def _master_fold_index(master: dict[str, dict[str, Any]]) -> dict[str, str]:
    index: dict[str, str] = {}
    for artist_id, entry in master.items():
        for alias in [entry.get("display_name"), *(entry.get("aliases") or [])]:
            fold = artist_fold(alias)
            if fold:
                index[fold] = artist_id
    return index


def _possible_matches(group: dict[str, Any], known: dict[str, str]) -> list[dict[str, str]]:
    matches: dict[tuple[str, str], dict[str, str]] = {}
    for name in group["names"]:
        stripped = artist_fold(strip_biography(name))
        if stripped and stripped in known:
            artist_id = known[stripped]
            matches[(artist_id, "strip_biography")] = {
                "artist_id": artist_id,
                "method": "strip_biography",
                "source_name": name,
            }
        inverted = propose_inversion(name)
        if inverted and inverted in known:
            artist_id = known[inverted]
            matches[(artist_id, "comma_inversion")] = {
                "artist_id": artist_id,
                "method": "comma_inversion",
                "source_name": name,
            }
    return sorted(matches.values(), key=lambda item: (item["artist_id"], item["method"]))


def _field_status(canonical: Any, candidates: Counter) -> str:
    if canonical is not None:
        return "curated"
    if len(candidates) > 1:
        return "conflict"
    if len(candidates) == 1:
        return "candidate_available"
    return "missing"


def _catalog_conflicts(entry: dict[str, Any], group: dict[str, Any]) -> list[str]:
    flags: list[str] = []
    country = normalize_country(entry.get("country_birth"))
    nationalities = {
        value
        for value in (normalize_country(item) for item in entry.get("nationalities") or [])
        if value
    }
    allowed_countries = nationalities | ({country} if country else set())
    observed_countries = set(group["catalog_countries"]) | set(group["bio_countries"])
    if allowed_countries and observed_countries - allowed_countries:
        flags.append("catalog_country_conflict")

    birth = _safe_year(entry.get("birth_year"))
    death = _safe_year(entry.get("death_year"))
    if birth and set(group["birth_years"]) - {birth}:
        flags.append("catalog_birth_year_conflict")
    if death and set(group["death_years"]) - {death}:
        flags.append("catalog_death_year_conflict")
    return flags


def build_dimension(groups: dict[str, dict[str, Any]]) -> list[dict[str, Any]]:
    master = load_master()
    known = _master_fold_index(master)
    # El maestro completo entra aunque no tenga lotes en el snapshot actual.
    for artist_id in master:
        groups.setdefault(artist_id, _new_group(artist_id, artist_id))

    rows: list[dict[str, Any]] = []
    for key, group in groups.items():
        artist_id = group["artist_id"]
        entry = master.get(artist_id) if artist_id else None
        aliases = [name for name, _ in group["names"].most_common()]
        possible_matches: list[dict[str, str]] = []

        if entry:
            display_name = entry.get("display_name") or (aliases[0] if aliases else artist_id)
            aliases = list(
                dict.fromkeys(
                    [
                        *(entry.get("aliases") or []),
                        entry.get("display_name"),
                        *aliases,
                    ]
                )
            )
            aliases = [str(value) for value in aliases if value]
            country_birth = normalize_country(entry.get("country_birth"))
            birth_year = _safe_year(entry.get("birth_year"))
            death_year = _safe_year(entry.get("death_year"))
            entity_type = entry.get("entity_type") or "artist"
            evidence_urls = _notes_urls(entry.get("notes"))
            missing_fields = [
                field
                for field, value in (
                    ("country_birth", country_birth),
                    ("birth_year", birth_year),
                    ("death_year", death_year),
                )
                if value is None
            ]
            quality_flags = _catalog_conflicts(entry, group)
            if not evidence_urls:
                quality_flags.append("no_evidence_url")
            if not any(attribution_type(alias) == "autor" for alias in aliases):
                quality_flags.append("no_resolvable_author_alias")
            if any(flag.startswith("catalog_") for flag in quality_flags):
                curation_status = "needs_review"
            elif missing_fields or not evidence_urls:
                curation_status = "needs_web"
            else:
                curation_status = "verified"
            confidence = entry.get("confidence")
            source = entry.get("source")
        else:
            display_name = aliases[0] if aliases else key
            country_birth = birth_year = death_year = None
            entity_type = "unknown"
            evidence_urls = []
            missing_fields = list(CANONICAL_FIELDS)
            quality_flags = ["identity_unverified"]
            possible_matches = _possible_matches(group, known)
            evidence_conflict = any(
                len(group[field]) > 1
                for field in ("bio_countries", "birth_years", "death_years")
            )
            if evidence_conflict:
                quality_flags.append("biographical_evidence_conflict")
            if possible_matches:
                quality_flags.append("possible_existing_master_alias")
            curation_status = "ambiguous" if evidence_conflict or possible_matches else "needs_web"
            confidence = "low"
            source = "observed"

        candidate_country = _single(group["bio_countries"])
        candidate_birth = _single(group["birth_years"])
        candidate_death = _single(group["death_years"])
        field_status = {
            "country_birth": _field_status(country_birth, group["bio_countries"]),
            "birth_year": _field_status(birth_year, group["birth_years"]),
            "death_year": _field_status(death_year, group["death_years"]),
        }
        death_status = entry.get("death_status") if entry else None
        if not death_status:
            death_status = "known" if death_year is not None else "unknown"

        rows.append(
            {
                "artist_key": key,
                "artist_id": artist_id,
                "display_name": display_name,
                "aliases": aliases,
                "curation_status": curation_status,
                "entity_type": entity_type,
                "confidence": confidence,
                "source": source,
                "country_birth": country_birth,
                "birth_year": birth_year,
                "death_year": death_year,
                "death_status": death_status,
                "candidate_country_birth": candidate_country,
                "candidate_birth_year": candidate_birth,
                "candidate_death_year": candidate_death,
                "catalog_country_candidates": _counter_rows(group["catalog_countries"]),
                "field_status": field_status,
                "missing_fields": missing_fields,
                "needs_web_research": curation_status != "verified",
                "quality_flags": sorted(set(quality_flags)),
                "possible_master_matches": possible_matches,
                "evidence_urls": evidence_urls,
                "evidence": {
                    "country_birth": _counter_rows(group["bio_countries"]),
                    "birth_year": _counter_rows(group["birth_years"]),
                    "death_year": _counter_rows(group["death_years"]),
                },
                "lots_observed": group["lots_observed"],
                "legacy_rows_observed": group["legacy_rows"],
                "houses": sorted(group["houses"]),
                "source_counts": {
                    name: count for name, count in sorted(group["sources"].items())
                },
                "source_files": sorted(group["source_files"]),
                "sample_lot_urls": group["lot_urls"],
            }
        )

    status_order = {name: i for i, name in enumerate(CURATION_STATUSES)}
    rows.sort(
        key=lambda row: (
            status_order.get(row["curation_status"], 99),
            -row["lots_observed"],
            row["artist_key"],
        )
    )
    return rows


def build_dispositions(
    dispositions: dict[str, dict[str, Any]],
    dimension: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    by_key = {row["artist_key"]: row for row in dimension}
    rows = []
    for item in dispositions.values():
        attribution = item["attribution_type"]
        dim = by_key.get(item["artist_key"])
        if attribution == "autor" and dim:
            disposition = (
                "master"
                if dim["artist_id"]
                else dim["curation_status"]
            )
        elif attribution == "desconocido":
            disposition = "unresolved"
        elif attribution == "no_autor":
            disposition = "rejected_non_artist"
        else:
            disposition = "excluded_attribution"
        rows.append(
            {
                "artist_fold": item["artist_fold"],
                "artist_key": item["artist_key"],
                "raw_names": [name for name, _ in item["names"].most_common()],
                "attribution_type": attribution,
                "disposition": disposition,
                "lots": item["lots"],
                "houses": sorted(item["houses"]),
                "source_counts": dict(sorted(item["sources"].items())),
            }
        )
    rows.sort(
        key=lambda row: (
            row["disposition"],
            -row["lots"],
            row["artist_fold"] or "",
        )
    )
    return rows


def _sha256(path: Path) -> str | None:
    if not path.exists():
        return None
    digest = hashlib.sha256()
    with open(path, "rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _source_path(path: Path) -> str:
    """Ruta legible; los tests y ejecuciones ad-hoc pueden vivir fuera del repo."""
    try:
        return str(path.relative_to(ROOT))
    except ValueError:
        return str(path)


def coverage_summary(
    dimension: list[dict[str, Any]],
    dispositions: list[dict[str, Any]],
    latent: list[dict[str, Any]],
    stats: dict[str, int],
    lots_path: Path,
) -> dict[str, Any]:
    author_dispositions = [
        row for row in dispositions if row["attribution_type"] == "autor"
    ]
    covered = [
        row
        for row in author_dispositions
        if row["disposition"] in {"master", "verified", "needs_web", "needs_review", "ambiguous"}
    ]
    field_counts = {
        field: sum(1 for row in dimension if row[field] is not None)
        for field in CANONICAL_FIELDS
    }
    statuses = Counter(row["curation_status"] for row in dimension)
    dispositions_count = Counter(row["disposition"] for row in dispositions)
    author_total = len(author_dispositions)
    return {
        "source": {
            "silver_lots": _source_path(lots_path),
            "silver_sha256": _sha256(lots_path),
            "legacy_excel_included": bool(stats.get("legacy_excel_present")),
        },
        "coverage_definition": (
            "Cada fold explicitamente atribuido a autor queda en maestro, needs_web "
            "o ambiguous; pendientes no se promueven automaticamente."
        ),
        "silver_lots": stats.get("silver_lots", 0),
        "author_lots": stats.get("author_lots", 0),
        "author_folds": author_total,
        "author_folds_accounted": len(covered),
        "author_fold_coverage_rate": round(len(covered) / author_total, 6)
        if author_total
        else 1.0,
        "dimension_rows": len(dimension),
        "curated_master_rows": sum(1 for row in dimension if row["artist_id"]),
        "pending_identity_rows": sum(1 for row in dimension if not row["artist_id"]),
        "curation_status_counts": dict(sorted(statuses.items())),
        "disposition_counts": dict(sorted(dispositions_count.items())),
        "canonical_field_counts": field_counts,
        "needs_web_research": sum(1 for row in dimension if row["needs_web_research"]),
        "rows_with_possible_master_match": sum(
            1 for row in dimension if row["possible_master_matches"]
        ),
        "rows_with_evidence_conflict": sum(
            1
            for row in dimension
            if "biographical_evidence_conflict" in row["quality_flags"]
        ),
        "latent_lefebre_candidates": len(latent),
        "latent_lefebre_lots": sum(row["lots"] for row in latent),
        "legacy_author_rows": stats.get("legacy_author_rows", 0),
    }


def write_jsonl(path: Path, rows: Iterable[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(path.suffix + ".tmp")
    with open(tmp, "w", encoding="utf-8", newline="\n") as handle:
        for row in rows:
            handle.write(json.dumps(row, ensure_ascii=False, sort_keys=True) + "\n")
    tmp.replace(path)


def write_json(path: Path, value: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(path.suffix + ".tmp")
    tmp.write_text(
        json.dumps(value, ensure_ascii=False, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
    )
    tmp.replace(path)


def write_review_queue(
    dimension: list[dict[str, Any]], summary: dict[str, Any]
) -> None:
    queue_fields = (
        "artist_key",
        "display_name",
        "aliases",
        "curation_status",
        "entity_type",
        "country_birth",
        "birth_year",
        "death_year",
        "death_status",
        "candidate_country_birth",
        "candidate_birth_year",
        "candidate_death_year",
        "catalog_country_candidates",
        "field_status",
        "missing_fields",
        "needs_web_research",
        "quality_flags",
        "possible_master_matches",
        "evidence",
        "lots_observed",
        "legacy_rows_observed",
        "houses",
        "source_counts",
        "sample_lot_urls",
    )
    pending = [
        {field: row[field] for field in queue_fields}
        for row in dimension
        if row["artist_id"] is None
    ]
    shards: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in pending:
        initial = row["artist_key"][:1].lower()
        shard = initial if "a" <= initial <= "z" else "_misc"
        shards[shard].append(row)
    REVIEW_DIR.mkdir(parents=True, exist_ok=True)
    for shard in [chr(code) for code in range(ord("a"), ord("z") + 1)] + ["_misc"]:
        write_jsonl(REVIEW_DIR / f"{shard}.jsonl", shards.get(shard, []))
    write_json(
        REVIEW_MANIFEST,
        {
            "schema_version": 1,
            "source": summary["source"],
            "review_queue_rows": len(pending),
            "shards": {
                shard: len(rows) for shard, rows in sorted(shards.items())
            },
            "curation_status_counts": dict(
                sorted(Counter(row["curation_status"] for row in pending).items())
            ),
            "rule": (
                "Registro exhaustivo de candidatos; no es cargado por load_master "
                "y no concede identidad ni pais."
            ),
        },
    )


def build(
    *,
    lots_path: Path = SILVER_LOTS,
    legacy_excel: Path | None = LEGACY_EXCEL,
    write_queue: bool = False,
) -> dict[str, Any]:
    groups, raw_dispositions, latent, silver_stats = collect_silver(lots_path)
    legacy_stats: dict[str, int] = {}
    if legacy_excel is not None:
        legacy_stats = collect_legacy_excel(legacy_excel, groups, raw_dispositions)
    stats = {**silver_stats, **legacy_stats}

    dimension = build_dimension(groups)
    dispositions = build_dispositions(raw_dispositions, dimension)
    summary = coverage_summary(dimension, dispositions, latent, stats, lots_path)
    summary["latent_candidates"] = latent

    write_jsonl(DIM_ARTISTS, dimension)
    write_jsonl(DISPOSITIONS, dispositions)
    write_json(COVERAGE, summary)
    if write_queue:
        write_review_queue(dimension, summary)
    return summary


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Construye el maestro operativo exhaustivo y la cola segura de revision."
    )
    parser.add_argument("--silver", type=Path, default=SILVER_LOTS)
    parser.add_argument(
        "--legacy-excel",
        type=Path,
        default=LEGACY_EXCEL,
        help="FINALL.xlsx opcional; si no existe se omite sin fallar.",
    )
    parser.add_argument(
        "--without-legacy-excel",
        action="store_true",
        help="No inspecciona FINALL.xlsx aunque exista.",
    )
    parser.add_argument(
        "--write-review-queue",
        action="store_true",
        help="Actualiza la cola versionada de candidatos pendientes.",
    )
    args = parser.parse_args()

    summary = build(
        lots_path=args.silver,
        legacy_excel=None if args.without_legacy_excel else args.legacy_excel,
        write_queue=args.write_review_queue,
    )
    print("--- Maestro operativo de artistas ---")
    for field in (
        "dimension_rows",
        "curated_master_rows",
        "pending_identity_rows",
        "author_folds",
        "author_folds_accounted",
        "author_fold_coverage_rate",
        "needs_web_research",
        "latent_lefebre_candidates",
    ):
        print(f"  {field}: {summary[field]}")
    print(f"[gold] wrote: {DIM_ARTISTS}")
    print(f"[gold] wrote: {DISPOSITIONS}")
    print(f"[gold] wrote: {COVERAGE}")
    if args.write_review_queue:
        print(f"[curation] wrote: {REVIEW_DIR}")


if __name__ == "__main__":
    main()
